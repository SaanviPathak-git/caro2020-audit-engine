"""
Deliverables Validation Tests (Excel, Markdown, PDF)
"""

import pytest
from pathlib import Path
import openpyxl
from caro_engine.core.engine import CaroAuditEngine
from caro_engine.deliverables.excel_workpaper import generate_excel_workpaper
from caro_engine.deliverables.draft_caro_report import save_draft_caro_report
from caro_engine.deliverables.pdf_report import generate_caro_pdf_report

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "data" / "sample_clients"

def test_deliverables_generation(tmp_path: Path):
    tata_path = SAMPLE_DIR / "tata_motors_fy24"
    engine = CaroAuditEngine(tata_path)
    summary = engine.run_audit()
    meta = summary.metadata
    
    excel_path = tmp_path / "Test_Workpaper.xlsx"
    md_path = tmp_path / "Test_Report.md"
    pdf_path = tmp_path / "Test_Report.pdf"
    
    # 1. Test Excel Workpaper
    generate_excel_workpaper(excel_path, meta, summary.clause_results)
    assert excel_path.exists()
    assert excel_path.stat().st_size > 5000  # Non-empty binary workbook
    wb = openpyxl.load_workbook(excel_path)
    assert "Lead Sheet & Materiality" in wb.sheetnames
    assert "CARO Clause Matrix" in wb.sheetnames
    assert "Exceptions Log" in wb.sheetnames
    
    # 2. Test Markdown Report
    save_draft_caro_report(md_path, meta, summary.clause_results)
    assert md_path.exists()
    content = md_path.read_text(encoding="utf-8")
    assert "ANNEXURE 'A' TO THE INDEPENDENT AUDITOR'S REPORT" in content
    assert "Tata Motors Limited" in content
    assert "UDIN" in content
    
    # 3. Test PDF Deliverable
    generate_caro_pdf_report(pdf_path, meta, summary.clause_results)
    assert pdf_path.exists()
    assert pdf_path.stat().st_size > 5000
